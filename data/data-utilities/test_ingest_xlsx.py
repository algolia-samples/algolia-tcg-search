"""Run with: python -m unittest test_ingest_xlsx"""

import os
import tempfile
import unittest
from pathlib import Path

# ingest.py exits at import without an event id; no Algolia or network access is used here.
os.environ.setdefault("ALGOLIA_EVENT_ID", "test-event")

import openpyxl  # noqa: E402

import ingest  # noqa: E402


def write_workbook(path: Path, rows: list) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test Set"
    ws.append(["Pokemon Name", "Number", "Card Type", "# in Machine", "Estimated Value"])
    for row in rows:
        ws.append(row)
    wb.save(path)


class ProcessXlsxFile(unittest.TestCase):
    def setUp(self):
        self.dir = Path(tempfile.mkdtemp())

    def run_file(self, path):
        records = []
        complete = ingest.process_xlsx_file(path, None, "idx", enrich=False, dump=records)
        return complete, records

    def test_readable_workbook_is_complete(self):
        path = self.dir / "good.xlsx"
        write_workbook(path, [["Pikachu", "25", "Rare", 2, 1.5]])
        complete, records = self.run_file(path)
        self.assertIs(complete, True)
        self.assertEqual([r["pokemon_name"] for r in records], ["Pikachu"])

    def test_unreadable_workbook_is_reported_not_swallowed(self):
        # The Gartner workbook once arrived corrupt ("Bad magic number") and --merge
        # carried on with only the other file's cards.
        path = self.dir / "corrupt.xlsx"
        path.write_bytes(b"this is not a zip file")
        complete, records = self.run_file(path)
        self.assertIs(complete, False)
        self.assertEqual(records, [])

    def test_rows_not_in_the_machine_are_not_a_failure(self):
        path = self.dir / "partial.xlsx"
        write_workbook(path, [["Pikachu", "25", "Rare", 2, 1.5], ["Mew", "151", "Rare", None, 9.0]])
        complete, records = self.run_file(path)
        self.assertIs(complete, True)
        self.assertEqual(len(records), 1)

    def test_card_sheet_missing_required_columns_is_incomplete(self):
        # Its cards would be silently omitted from a merge.
        path = self.dir / "malformed.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Broken Set"
        ws.append(["Pokemon Name", "Number", "# in Machine"])  # no Card Type / Estimated Value
        ws.append(["Pikachu", "25", 2])
        wb.save(path)
        complete, records = self.run_file(path)
        self.assertIs(complete, False)
        self.assertEqual(records, [])

    def test_workbook_with_no_card_sheet_is_incomplete(self):
        path = self.dir / "no-cards.xlsx"
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Site URL"
        ws.append(["https://example.com"])
        wb.save(path)
        complete, _ = self.run_file(path)
        self.assertIs(complete, False)

    def test_extra_non_card_sheet_is_not_a_failure(self):
        path = self.dir / "mixed.xlsx"
        write_workbook(path, [["Pikachu", "25", "Rare", 2, 1.5]])
        wb = openpyxl.load_workbook(path)
        wb.create_sheet("Site URL").append(["https://example.com"])
        wb.save(path)
        complete, records = self.run_file(path)
        self.assertIs(complete, True)
        self.assertEqual(len(records), 1)

    def test_hidden_sheets_do_not_count_as_card_sheets(self):
        path = self.dir / "hidden-only.xlsx"
        write_workbook(path, [["Pikachu", "25", "Rare", 2, 1.5]])
        wb = openpyxl.load_workbook(path)
        wb["Test Set"].sheet_state = "hidden"
        wb.create_sheet("Notes").append(["nothing here"])
        wb.save(path)
        complete, records = self.run_file(path)
        self.assertIs(complete, False)
        self.assertEqual(records, [])


if __name__ == "__main__":
    unittest.main()
